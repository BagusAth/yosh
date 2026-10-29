#!/usr/bin/env python3
"""
Audit label manual 150 citra (RQ4, Pakem v1.3.1, Fase 5)
========================================================
Desain (pakem, bagian Analisis): 100 citra latih dengan D_image tertinggi + 50 citra acak sebagai pembanding.
Tambahan agar audit sahih:
  - Maksimal satu citra per grup foto (salinan augmentasi tidak dihitung dua kali).
  - Buta strata: urutan diacak, kode A001..A150, strata dan skor disimpan terpisah di _kunci/.
  - 30 citra (20 tinggi + 10 acak) untuk rater kedua, di file terpisah agar tidak saling melihat.

Keluaran (folder baru, tidak menyentuh dataset):
  audit_label_150/gambar/Annn.jpg            kiri citra asli, kanan kotak label bernomor
  audit_label_150/LEMBAR_AUDIT_LABEL.xlsx     diisi rater 1 (tidak ditimpa bila sudah ada)
  audit_label_150/LEMBAR_AUDIT_RATER2.xlsx    diisi rater 2 (opsional)
  audit_label_150/_kunci/audit_key.csv        strata, peringkat, D_image, skor kualitas label (buka setelah audit selesai)

Pemakaian: python build_label_audit.py [--root .] [--overwrite-sheets]
"""
from __future__ import annotations

import argparse
import csv
import random
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont

CLASS_NAMES = ["CBB", "CBSD", "CGM", "CMD", "HEALTHY"]
CLASS_COLORS = {0: (240, 228, 66), 1: (204, 121, 167), 2: (86, 180, 233), 3: (213, 94, 0), 4: (255, 255, 255)}
N_TOP, N_RAND, N_R2_TOP, N_R2_RAND = 100, 50, 20, 10
SEED_RAND, SEED_ORDER, SEED_R2 = 42, 2026, 7

KATEGORI = [
    ("Anotasi tidak lengkap", "Ada daun bergejala (atau daun sehat pada citra HEALTHY) yang jelas terlihat tetapi tidak diberi kotak."),
    ("Kotak bersarang/ganda", "Kotak berada di dalam kotak lain, atau satu objek diberi lebih dari satu kotak."),
    ("Kelas ambigu/salah", "Gejala bisa masuk lebih dari satu kelas, atau kelas pada label tampak keliru."),
    ("Objek bukan daun", "Kotak menandai batang, umbi, atau bagian tanaman lain selain daun (misalnya lesi CBSD pada batang)."),
    ("Kotak tidak presisi", "Objek dan kelasnya benar, tetapi kotak terlalu longgar, terlalu sempit, atau bergeser."),
    ("Artefak citra", "Citra buram, terlalu gelap/terang, tertutup bayangan, terpotong, atau ada isian hitam hasil rotasi yang mengganggu."),
    ("Benar-benar sulit", "Label tampak benar dan lengkap, tetapi gejala kecil, samar, atau tertutup kanopi padat."),
    ("Tidak ada masalah", "Label benar, lengkap, dan citra tidak tampak sulit."),
    ("Lainnya", "Tidak cocok dengan kategori di atas; jelaskan di kolom Catatan."),
]
TAK_BERLABEL = ["0", "1-2", "3-5", ">5"]
YA_TIDAK = ["Ya", "Sebagian", "Tidak", "Ragu"]

COLS = [  # (judul, lebar, diisi?)
    ("No", 5, False), ("Kode", 8, False), ("Gambar", 11, False), ("Jumlah kotak", 9, False), ("Kelas pada label", 24, False),
    ("Kategori utama", 24, True), ("Kategori tambahan (opsional)", 24, True), ("Objek bergejala tak berlabel (perkiraan)", 15, True),
    ("Kelas label sesuai?", 12, True), ("Kotak presisi?", 12, True), ("Catatan (sebut nomor kotak bila perlu)", 42, True),
]

ARIAL = "Arial"
F_BODY, F_HEAD = Font(name=ARIAL, size=10), Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
FILL_HEAD, FILL_IN = PatternFill("solid", fgColor="2F5597"), PatternFill("solid", fgColor="FFF2CC")
FILL_DONE = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def load_groups(root: Path) -> dict[str, str]:
    m = pd.read_csv(root / "cassava-v7-stage-6" / "resplit_group_manifest.csv", encoding="utf-8-sig")
    g = {}
    for gid, mem in zip(m.clean_group_id, m.member_filenames):
        for fn in str(mem).split(";"):
            if fn.strip():
                g[fn.strip()] = str(gid)
    return g


def read_labels(path: Path) -> list[tuple[int, float, float, float, float]]:
    if not path.exists():
        return []
    out = []
    for line in path.read_text(encoding="utf-8").splitlines():
        p = line.split()
        if len(p) >= 5:
            out.append((int(float(p[0])), *map(float, p[1:5])))
    return out


def get_font(size: int):
    for name in ("arialbd.ttf", "arial.ttf", "DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def render(img_path: Path, labels, code: str, out: Path) -> None:
    im = Image.open(img_path).convert("RGB")
    w, h = im.size
    head = 34
    canvas = Image.new("RGB", (w * 2 + 10, h + head), (30, 30, 30))
    canvas.paste(im, (0, head))
    ann = im.copy()
    d = ImageDraw.Draw(ann)
    f = get_font(15)
    for i, (c, cx, cy, bw, bh) in enumerate(labels, start=1):
        x1, y1, x2, y2 = (cx - bw / 2) * w, (cy - bh / 2) * h, (cx + bw / 2) * w, (cy + bh / 2) * h
        col = CLASS_COLORS.get(c, (255, 0, 255))
        d.rectangle([x1, y1, x2, y2], outline=(0, 0, 0), width=5)
        d.rectangle([x1, y1, x2, y2], outline=col, width=3)
        txt = f"{i} {CLASS_NAMES[c] if 0 <= c < len(CLASS_NAMES) else c}"
        tb = d.textbbox((0, 0), txt, font=f)
        tw, th = tb[2] - tb[0], tb[3] - tb[1]
        ty = y1 - th - 6 if y1 - th - 6 > 0 else y1 + 2
        d.rectangle([x1, ty, x1 + tw + 6, ty + th + 6], fill=col)
        d.text((x1 + 3, ty + 1), txt, font=f, fill=(0, 0, 0))
    canvas.paste(ann, (w + 10, head))
    dh = ImageDraw.Draw(canvas)
    fh = get_font(18)
    dh.text((8, 7), f"{code}   |   kiri: citra asli   |   kanan: label ({len(labels)} kotak)", font=fh, fill=(255, 255, 255))
    x = w * 2 + 10 - 8
    for c in reversed(range(len(CLASS_NAMES))):
        tb = dh.textbbox((0, 0), CLASS_NAMES[c], font=fh)
        x -= tb[2] - tb[0] + 26
        dh.rectangle([x, 9, x + 14, 23], fill=CLASS_COLORS[c])
        dh.text((x + 18, 7), CLASS_NAMES[c], font=fh, fill=(255, 255, 255))
    out.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(out, quality=90)


def style_header(ws, row: int, ncol: int) -> None:
    for j in range(1, ncol + 1):
        c = ws.cell(row=row, column=j)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def add_lists_sheet(wb) -> dict[str, str]:
    ws = wb.create_sheet("Daftar")
    ranges = {}
    for j, (name, vals) in enumerate((("Kategori", [k for k, _ in KATEGORI]), ("TakBerlabel", TAK_BERLABEL), ("YaTidak", YA_TIDAK)), start=1):
        ws.cell(row=1, column=j, value=name).font = Font(name=ARIAL, bold=True)
        for i, v in enumerate(vals, start=2):
            ws.cell(row=i, column=j, value=v).font = F_BODY
        col = get_column_letter(j)
        ranges[name] = f"=Daftar!${col}$2:${col}${len(vals) + 1}"
    ws.sheet_state = "hidden"
    return ranges


def add_audit_sheet(wb, title: str, rows: list[dict], ranges: dict[str, str]):
    ws = wb.create_sheet(title)
    for j, (name, width, _) in enumerate(COLS, start=1):
        ws.cell(row=1, column=j, value=name)
        ws.column_dimensions[get_column_letter(j)].width = width
    style_header(ws, 1, len(COLS))
    ws.row_dimensions[1].height = 42
    for i, r in enumerate(rows, start=2):
        vals = [i - 1, r["kode"], "Buka", r["n_gt"], r["kelas"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border, c.alignment = F_BODY, BORDER, WRAP
        link = ws.cell(row=i, column=3)
        link.hyperlink = f"gambar/{r['kode']}.jpg"
        link.font = Font(name=ARIAL, size=10, color="0563C1", underline="single")
        for j in range(6, len(COLS) + 1):
            c = ws.cell(row=i, column=j)
            c.fill, c.border, c.font, c.alignment = FILL_IN, BORDER, F_BODY, WRAP
    last = len(rows) + 1
    for col, key in (("F", "Kategori"), ("G", "Kategori"), ("H", "TakBerlabel"), ("I", "YaTidak"), ("J", "YaTidak")):
        dv = DataValidation(type="list", formula1=ranges[key], allow_blank=True, showErrorMessage=True,
                            errorTitle="Pilihan tidak dikenal", error="Pilih dari daftar (klik panah di sel).")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    ws.conditional_formatting.add(f"A2:K{last}", FormulaRule(formula=['$F2<>""'], fill=FILL_DONE))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:K{last}"
    return ws


def add_guide_sheet(wb, rater: str, n_rows: int, sheet_name: str) -> None:
    ws = wb.active
    ws.title = "Panduan"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 95
    lines = [
        (f"Audit label manual — {rater}", None),
        (f"Isi sheet \"{sheet_name}\" ({n_rows} citra). Sel kuning adalah sel yang diisi; baris berubah hijau setelah Kategori utama terisi.", None),
        ("Langkah", "1) Klik \"Buka\" di kolom Gambar (atau buka gambar/Annn.jpg). Kiri citra asli, kanan kotak label bernomor."),
        ("", "2) Bandingkan kedua panel. Cari daun bergejala tanpa kotak, kotak salah kelas, kotak tidak pas, atau kotak ganda."),
        ("", "3) Pilih Kategori utama (yang paling menjelaskan kualitas label atau kesulitan citra), lalu kolom lain bila relevan."),
        ("", "4) Jangan membuka folder _kunci sebelum semua citra selesai dinilai. Urutan citra sengaja diacak."),
        ("", "5) Simpan file secara berkala. Perkiraan waktu: 1–2 menit per citra."),
        ("Definisi kategori", None),
    ]
    r = 1
    for a, b in lines:
        ws.cell(row=r, column=1, value=a).font = Font(name=ARIAL, size=13 if r == 1 else 10, bold=(r == 1 or b is None and r > 2))
        if b is not None:
            ws.cell(row=r, column=2, value=b).font = F_BODY
            ws.cell(row=r, column=2).alignment = WRAP
        r += 1
    ws.cell(row=2, column=1).alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
    ws.row_dimensions[2].height = 30
    ws.cell(row=r, column=1, value="Kategori")
    ws.cell(row=r, column=2, value="Gunakan bila")
    style_header(ws, r, 2)
    r += 1
    for k, desc in KATEGORI:
        for j, v in enumerate((k, desc), start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border, c.alignment = F_BODY, BORDER, WRAP
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Kolom lain").font = Font(name=ARIAL, size=10, bold=True)
    r += 1
    for a, b in (
        ("Objek bergejala tak berlabel", "Perkiraan jumlah objek yang seharusnya diberi kotak tetapi tidak: 0, 1-2, 3-5, >5."),
        ("Kelas label sesuai?", "Apakah kelas pada kotak-kotak yang ada sesuai gejala: Ya, Sebagian, Tidak, Ragu."),
        ("Kotak presisi?", "Apakah batas kotak pas dengan objeknya: Ya, Sebagian, Tidak, Ragu."),
    ):
        ws.cell(row=r, column=1, value=a).font = F_BODY
        ws.cell(row=r, column=2, value=b).font = F_BODY
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Contoh pengisian (bukan data)").font = Font(name=ARIAL, size=10, bold=True)
    r += 1
    ex = ["Anotasi tidak lengkap", "Benar-benar sulit", "3-5", "Ya", "Sebagian", "Hanya 2 dari ~6 daun bergejala dilabeli; kotak 2 terlalu longgar."]
    heads = [COLS[j][0] for j in range(5, len(COLS))]
    for j, (hd, v) in enumerate(zip(heads, ex)):
        ws.cell(row=r + j, column=1, value=hd).font = F_BODY
        c = ws.cell(row=r + j, column=2, value=v)
        c.font, c.fill = Font(name=ARIAL, size=10, italic=True), FILL_IN


def add_progress_sheet(wb, sheet_name: str, n_rows: int) -> None:
    ws = wb.create_sheet("Progres")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    last = n_rows + 1
    ws["A1"], ws["B1"] = "Ringkasan pengisian", "Jumlah"
    style_header(ws, 1, 2)
    ws["A2"], ws["B2"] = "Citra dinilai", f"=COUNTA('{sheet_name}'!F2:F{last})"
    ws["A3"], ws["B3"] = "Citra total", n_rows
    ws["A4"], ws["B4"] = "Selesai", "=IF(B3=0,0,B2/B3)"
    ws["B4"].number_format = "0%"
    ws["A6"], ws["B6"] = "Kategori utama", "Jumlah"
    style_header(ws, 6, 2)
    for i, (k, _) in enumerate(KATEGORI, start=7):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=f"=COUNTIF('{sheet_name}'!F2:F{last},A{i})")
    for row in ws.iter_rows(min_row=2, max_row=6 + len(KATEGORI)):
        for c in row:
            if c.font is None or not c.font.bold:
                c.font = F_BODY
    ws["A16"] = "Hitungan per strata baru dibuat setelah audit selesai (lihat _kunci/audit_key.csv)."
    ws["A16"].font = Font(name=ARIAL, size=9, italic=True)


def build_workbook(path: Path, rater: str, sheet_name: str, rows: list[dict]) -> None:
    wb = Workbook()
    add_guide_sheet(wb, rater, len(rows), sheet_name)
    ranges = add_lists_sheet(wb)
    add_audit_sheet(wb, sheet_name, rows, ranges)
    add_progress_sheet(wb, sheet_name, len(rows))
    wb.move_sheet("Daftar", offset=10)
    wb.active = 1
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".")
    ap.add_argument("--out", default="audit_label_150")
    ap.add_argument("--overwrite-sheets", action="store_true", help="Timpa lembar xlsx yang sudah ada (HATI-HATI: isian hilang)")
    args = ap.parse_args()
    root = Path(args.root).resolve()
    out = root / args.out
    diff_dir = root / "stage_v2_difficulty_20261028"
    man = pd.read_csv(diff_dir / "train_image_difficulty_manifest.csv")
    lq = pd.read_csv(diff_dir / "train_image_label_quality.csv")
    df = man.merge(lq[["image_id", "n_gt", "n_suspect_missing", "n_suspect_swap", "n_suspect_badloc", "label_quality"]], on="image_id", how="left")
    groups = load_groups(root)
    df["group_id"] = df.image_id.map(lambda n: groups.get(n, f"tanpa_grup:{n}"))
    # 112 citra punya D_image = 1,0 persis: urutan seri diacak (bukan urutan nama file) sebelum diurutkan
    df = df.sample(frac=1.0, random_state=SEED_RAND).sort_values("d_image", ascending=False, kind="mergesort").reset_index(drop=True)
    df["rank_d"] = range(1, len(df) + 1)

    taken_groups, picks = set(), []
    for r in df.itertuples():
        if len(picks) >= N_TOP:
            break
        if r.group_id not in taken_groups:
            taken_groups.add(r.group_id)
            picks.append((r.Index, "tinggi"))
    rest = [i for i in df.index if df.at[i, "group_id"] not in taken_groups]
    random.Random(SEED_RAND).shuffle(rest)
    for i in rest:
        if sum(1 for _, s in picks if s == "acak") >= N_RAND:
            break
        g = df.at[i, "group_id"]
        if g not in taken_groups:
            taken_groups.add(g)
            picks.append((i, "acak"))
    random.Random(SEED_ORDER).shuffle(picks)

    img_dir = root / "cassava-v7-runtime-mirror" / "train" / "images"
    lbl_dir = root / "cassava-v7-runtime-mirror" / "train" / "labels"
    rows = []
    for k, (i, strata) in enumerate(picks, start=1):
        r = df.loc[i]
        code = f"A{k:03d}"
        labels = read_labels(lbl_dir / (Path(r.image_id).stem + ".txt"))
        cnt = Counter(CLASS_NAMES[c] if 0 <= c < len(CLASS_NAMES) else str(c) for c, *_ in labels)
        rows.append({"kode": code, "image_id": r.image_id, "strata": strata, "rank_d": int(r.rank_d), "d_image": float(r.d_image),
                     "label_quality": r.label_quality, "n_suspect_missing": r.n_suspect_missing, "n_suspect_swap": r.n_suspect_swap,
                     "n_suspect_badloc": r.n_suspect_badloc, "n_gt": len(labels), "group_id": r.group_id,
                     "kelas": ", ".join(f"{c} x{n}" for c, n in sorted(cnt.items())) or "(tanpa kotak)", "_labels": labels})
    r2_top = [x for x in rows if x["strata"] == "tinggi"]
    r2_rand = [x for x in rows if x["strata"] == "acak"]
    rng = random.Random(SEED_R2)
    r2 = set(x["kode"] for x in rng.sample(r2_top, N_R2_TOP) + rng.sample(r2_rand, N_R2_RAND))
    for x in rows:
        x["rater2"] = int(x["kode"] in r2)

    n_new = 0
    for x in rows:
        dst = out / "gambar" / f"{x['kode']}.jpg"
        render(img_dir / x["image_id"], x["_labels"], x["kode"], dst)  # deterministik: aman ditulis ulang
        n_new += 1

    key_dir = out / "_kunci"
    key_dir.mkdir(parents=True, exist_ok=True)
    keys = ["kode", "image_id", "strata", "rank_d", "d_image", "label_quality", "n_suspect_missing", "n_suspect_swap",
            "n_suspect_badloc", "n_gt", "group_id", "rater2"]
    with open(key_dir / "audit_key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    (key_dir / "BACA_DULU.txt").write_text(
        "Jangan buka audit_key.csv sebelum semua citra selesai dinilai: file ini berisi strata (tinggi/acak) dan skor kesulitan.\n",
        encoding="utf-8")

    for fname, rater, sheet, sub in (("LEMBAR_AUDIT_LABEL.xlsx", "rater 1", "Audit", rows),
                                     ("LEMBAR_AUDIT_RATER2.xlsx", "rater 2 (opsional, 30 citra)", "Rater2",
                                      [x for x in rows if x["rater2"]])):
        p = out / fname
        if p.exists() and not args.overwrite_sheets:
            print(f"{fname} sudah ada; tidak ditimpa.")
            continue
        build_workbook(p, rater, sheet, sub)
    s = Counter(x["strata"] for x in rows)
    print(f"Selesai: {len(rows)} citra ({dict(s)}), {n_new} gambar ditulis, rater 2: {sum(x['rater2'] for x in rows)} citra -> {out}")


if __name__ == "__main__":
    main()
